"""
Send provisional marksheet PDFs to students via Gmail.
Usage:
  python send_emails.py \
    --pdf-dir "sem6_provisional" \
    --excel "StudentData.xlsx" \
    --email-col "E Mail ID" \
    --sender "you@gmail.com" \
    --password "xxxx xxxx xxxx xxxx"
"""

import argparse
import os
import re
import sys
import time
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
import openpyxl


def load_email_map(excel_path: str, email_col: str, uid_col: str = 'UIDNumber') -> dict:
    """Return {enrollment_no: email} from Excel."""
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    headers = [str(c.value).strip() if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
    try:
        uid_idx = headers.index(uid_col)
    except ValueError:
        # Try case-insensitive
        uid_idx = next((i for i, h in enumerate(headers) if h.lower() == uid_col.lower()), None)
        if uid_idx is None:
            raise ValueError(f"Column '{uid_col}' not found. Headers: {headers}")
    try:
        email_idx = headers.index(email_col)
    except ValueError:
        email_idx = next((i for i, h in enumerate(headers) if h.lower() == email_col.lower()), None)
        if email_idx is None:
            raise ValueError(f"Column '{email_col}' not found. Headers: {headers}")

    mapping = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        uid = str(row[uid_idx]).strip() if row[uid_idx] else ''
        email = str(row[email_idx]).strip() if row[email_idx] else ''
        if uid and email and '@' in email:
            mapping[uid] = email

    print(f"Loaded {len(mapping)} student email records from Excel")
    return mapping


def send_email(sender: str, password: str, recipient: str, student_name: str, pdf_path: str):
    msg = MIMEMultipart()
    msg['From'] = sender
    msg['To'] = recipient
    msg['Subject'] = f'Provisional Marksheet — {student_name}'

    body = f"""Dear {student_name},

Please find attached your Provisional Marksheet.

This is an auto-generated email. For any queries, contact your institute.

Regards,
Examination Cell"""
    msg.attach(MIMEText(body, 'plain'))

    with open(pdf_path, 'rb') as f:
        part = MIMEBase('application', 'octet-stream')
        part.set_payload(f.read())
    encoders.encode_base64(part)
    part.add_header('Content-Disposition', f'attachment; filename="{os.path.basename(pdf_path)}"')
    msg.attach(part)

    with smtplib.SMTP_SSL('smtp.gmail.com', 465) as server:
        server.login(sender, password)
        server.sendmail(sender, recipient, msg.as_string())


def main():
    parser = argparse.ArgumentParser(description='Send provisional marksheets via Gmail')
    parser.add_argument('--pdf-dir', required=True)
    parser.add_argument('--excel', required=True)
    parser.add_argument('--email-col', required=True)
    parser.add_argument('--sender', required=True)
    parser.add_argument('--password', required=True)
    parser.add_argument('--uid-col', default='UIDNumber')
    parser.add_argument('--delay', type=float, default=3.0, help='Seconds between emails')
    args = parser.parse_args()

    email_map = load_email_map(args.excel, args.email_col, args.uid_col)

    pdfs = [f for f in os.listdir(args.pdf_dir) if f.endswith('.pdf')]
    print(f"Found {len(pdfs)} PDFs in {args.pdf_dir}")

    log_path = os.path.join(args.pdf_dir, 'email_log.txt')
    sent = 0
    failed = 0

    with open(log_path, 'w', encoding='utf-8') as log:
        log.write(f"Email log\n{'='*60}\n")
        for pdf_file in pdfs:
            # Extract enrollment from filename: "Name - 1234567890.pdf"
            m = re.search(r'-\s*(\d{10})\.pdf$', pdf_file)
            if not m:
                msg = f"SKIP {pdf_file} — enrollment not in filename"
                print(msg)
                log.write(msg + '\n')
                continue

            enrollment = m.group(1)
            email = email_map.get(enrollment)
            if not email:
                msg = f"SKIP {pdf_file} — no email for enrollment {enrollment}"
                print(msg)
                log.write(msg + '\n')
                failed += 1
                continue

            # Extract student name from filename
            student_name = pdf_file.rsplit('-', 1)[0].strip()
            pdf_path = os.path.join(args.pdf_dir, pdf_file)

            try:
                send_email(args.sender, args.password, email, student_name, pdf_path)
                msg = f"SENT {pdf_file} → {email}"
                print(msg)
                log.write(msg + '\n')
                sent += 1
                time.sleep(args.delay)
            except Exception as e:
                msg = f"FAIL {pdf_file} → {email} — {e}"
                print(msg, file=sys.stderr)
                log.write(msg + '\n')
                failed += 1

        summary = f"\nDone. Sent: {sent}  Failed: {failed}"
        print(summary)
        log.write(summary + '\n')

    print(f"Log saved: {log_path}")


if __name__ == '__main__':
    main()
